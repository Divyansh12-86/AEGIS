#!/usr/bin/env python3
"""Build the complete Schneider Electric PredictX pitch deck (.pptx).

Takes the official template 'PredictX Team Pitch Deck Template V2.pptx',
populates all slides with AEGIS technical rigor, embeds high-res figures,
adds a dedicated 'Honest Engineering & Limitations' slide (Priority 3), and appends
the two mandatory kickoff slides (Tagline Slide 13 & Closing Slide 14 from Kickoff PDF).

Outputs:
    PredictX_AEGIS_Final_Pitch_Deck.pptx (10 Slides)
"""
import copy
import sys
from pathlib import Path
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

ROOT = Path(__file__).resolve().parent.parent
TEMPLATE_PATH = ROOT / "PredictX Team Pitch Deck Template V2.pptx"
OUTPUT_PATH = ROOT / "PredictX_AEGIS_Final_Pitch_Deck.pptx"
FIG_TRAJECTORY = ROOT / "experiments/figures/degradation_trajectories/DS01_u2_trajectory.png"
IMG_SLIDE13 = ROOT / "assets/schneider_slide13_tagline.png"
IMG_SLIDE14 = ROOT / "assets/schneider_slide14_closing.png"


def set_text(shape, text, font_size=None, bold=None, color=None):
    """Safely replace text in shape preserving formatting where possible."""
    if not shape.has_text_frame:
        return
    tf = shape.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = text
    if font_size:
        p.font.size = Pt(font_size)
    if bold is not None:
        p.font.bold = bold
    if color:
        p.font.color.rgb = color


def main():
    print(f"Loading template from {TEMPLATE_PATH.name}...")
    prs = pptx.Presentation(TEMPLATE_PATH)

    # =========================================================================
    # SLIDE 1: Cover Slide
    # =========================================================================
    s1 = prs.slides[0]
    for sp in list(s1.shapes):
        if sp.shape_id in [9, 10, 13, 14, 15, 16, 25, 28, 29]:
            sp._element.getparent().remove(sp._element)

    for shape in s1.shapes:
        if shape.shape_id == 8 and shape.has_text_frame:
            set_text(shape, "CHOSEN HACKATHON THEME", font_size=11, bold=True, color=RGBColor(120, 120, 120))
        elif shape.shape_id == 11:
            shape.left = 777240
            shape.width = 5800000
        elif shape.shape_id == 12 and shape.has_text_frame:
            shape.left = 870000
            shape.width = 5600000
            set_text(shape, "✔ Energy Intelligence & Sustainability (Predictive Maintenance)", font_size=12, bold=True, color=RGBColor(0, 160, 50))
        elif shape.has_text_frame:
            t = shape.text_frame.text.strip()
            if "[ Your Team Name ]" in t:
                set_text(shape, "AEGIS — Event Grammar Probabilistic Model (EGPM)", font_size=24, bold=True, color=RGBColor(0, 160, 50))
            elif "[ Branch Name ]" in t:
                set_text(shape, "Computer Science & Electrical Engineering (Interdisciplinary Duo)", font_size=12, bold=True)
            elif "[ 2-4 Members ]" in t:
                set_text(shape, "2 Members: Atharva Mendhulkar & Divyansh Seth", font_size=12, bold=True)

    # =========================================================================
    # SLIDE 2: Idea Title, Problem, Solution, Tech Stack
    # =========================================================================
    s2 = prs.slides[1]
    for shape in s2.shapes:
        if shape.has_text_frame:
            t = shape.text_frame.text.strip()
            if "Idea Title" in t:
                set_text(shape, "AEGIS: Unified Neuro-Symbolic Predictive Maintenance", font_size=20, bold=True, color=RGBColor(0, 160, 50))
            elif t.startswith("[ In 2-3 sentences"):
                prob_text = (
                    "Critical electrical switchgear and data center cooling assets suffer from "
                    "architectural fragmentation: facilities run 4 disconnected black-box models "
                    "(RUL, anomaly, fault, health) that quadruplicate compute and drift maintenance. "
                    "Crucially, black boxes output uncalibrated point estimates and post-hoc SHAP illusions "
                    "that violate physical degradation wear physics, leaving reliability engineers unable to "
                    "certify safe operation."
                )
                set_text(shape, prob_text, font_size=11)
            elif t.startswith("[ Describe your solution in detail"):
                sol_text = (
                    "AEGIS unifies all predictive maintenance into a single inspectable state representation. "
                    "Multivariate telemetry is mapped via 1D causal convolutions and Vector Quantization (K=16) "
                    "into learned discrete events. An explicit-duration HSMM (M=5, D_max=10) models non-exponential "
                    "dwell physics. Four bottlenecked heads deliver Remaining Useful Life with 90% Conformal Prediction intervals, "
                    "acoustic anomaly detection, fault recognition, and health scoring with zero-hallucination "
                    "explanations and counterfactual what-if simulations."
                )
                set_text(shape, sol_text, font_size=11)
            elif t.startswith("[ Explain your technical approach"):
                tech_text = (
                    "Evaluated under locked causal Leave-One-Engine-Out (LOEO) protocol: achieves 18.1 RMSE on "
                    "N-CMAPSS DS01, outperforming deep CNN (34.0) and LSTM (38.9) by 15.9–20.8 RMSE (47–53% error reduction). "
                    "On MIMII industrial fans across 4 machines and 3 seeds, our continuous HSMM achieves 0.628 Macro AUROC. "
                    "Engine-level conformal calibration delivers verified 92% empirical coverage at 90% nominal confidence. "
                    "Validated by 260 automated tests passing with zero failures."
                )
                set_text(shape, tech_text, font_size=11)
            elif t.startswith("[ List languages, frameworks"):
                stack_text = (
                    "Python 3.14, PyTorch, NumPy, SciPy (Forward-Backward/Viterbi), Scikit-Learn, HDF5, Matplotlib, "
                    "Causal 1D-CNN, VQ-VAE (Corrected STE), Explicit-Duration HSMM, LOEO Conformal Calibration, PyTest (260 Tests). "
                    "Target Hardware: Schneider Harmony edge iPC & EcoStruxure Asset Advisor."
                )
                set_text(shape, stack_text, font_size=10)

    # =========================================================================
    # SLIDE 3: Innovation & Differentiation (25%)
    # =========================================================================
    s3 = prs.slides[2]
    for shape in s3.shapes:
        if shape.has_text_frame:
            t = shape.text_frame.text.strip()
            if t.startswith("[ How is this problem typically solved today"):
                today_text = (
                    "• 4 Disconnected Black Boxes: Separate independent models for RUL, acoustics, fault, and health quadrupling edge compute.\n"
                    "• Memoryless Assumption Flaw: Markov chains and geometric HMMs assume memoryless wear (constant hazard rates), violating mechanical and thermodynamic degradation physics.\n"
                    "• Uncalibrated Point Estimates: Models predict '35 cycles left' without error bounds, causing operators to scrap 25% of healthy asset life or face catastrophic arc-flash.\n"
                    "• Post-Hoc Illusions: SHAP/LIME heatmap attributions fluctuate wildly and lack physical causal validity."
                )
                set_text(shape, today_text, font_size=11)
            elif t.startswith("[ What specifically is different"):
                your_text = (
                    "1. Unified Neuro-Symbolic Grammar: 1 shared representation (discrete tokens + HSMM state posterior) powers all 4 tasks; cuts model size to <1.4 MB and edge compute TCO by 70%.\n"
                    "2. Explicit-Duration Degradation Physics: Non-parametric duration distributions p_m(d) model true cumulative wear; beats deep CNN/LSTM baselines by 15.9+ RMSE.\n"
                    "3. Rigorous LOEO Conformal Uncertainty: Engine-level 90% prediction intervals guarantee a safe lower bound for scheduled shutdowns with 92% verified empirical coverage.\n"
                    "4. Intrinsic Explanations & Counterfactuals: Deterministic root-cause objects with verified faithfulness (score drops 0.18–0.72) + counterfactual dwell-reset simulation."
                )
                set_text(shape, your_text, font_size=11)

    # =========================================================================
    # SLIDE 4: Prototype & Demo (20%)
    # =========================================================================
    s4 = prs.slides[3]
    for shape in list(s4.shapes):
        if shape.has_text_frame and "[ Replace this box with a screenshot" in shape.text_frame.text:
            shape.text_frame.clear()
        if shape.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE and shape.left < Inches(6):
            sp = shape._element
            sp.getparent().remove(sp)

    if FIG_TRAJECTORY.exists():
        print(f"Embedding prototype trajectory figure: {FIG_TRAJECTORY.name}")
        s4.shapes.add_picture(
            str(FIG_TRAJECTORY),
            left=Inches(0.75),
            top=Inches(1.95),
            width=Inches(6.9),
            height=Inches(4.1),
        )

    for shape in s4.shapes:
        if shape.has_text_frame:
            t = shape.text_frame.text.strip()
            if t.startswith("What stage is your prototype at?"):
                stage_text = (
                    "Stage: Working Production Prototype\n"
                    "Complete modular Python engine with 260/260 automated tests passing in CI. "
                    "Strict zero-leakage causal execution verified across all preprocessing and decoders."
                )
                set_text(shape, stage_text, font_size=10)
            elif t.startswith("What can it actually do today, live?"):
                live_text = (
                    "Live Capabilities (CLI & Python API):\n"
                    "Runs real-time sub-15ms edge inference: streams raw sensor telemetry, discovers discrete tokens (K=16), "
                    "infers HSMM state dwell, generates 4 unified outputs with 90% conformal intervals, and runs interactive "
                    "counterfactual what-if simulations via `scripts/demo_predictx.py`."
                )
                set_text(shape, live_text, font_size=10)
            elif t.startswith("What would the next version add?"):
                next_text = (
                    "Next Version Roadmap:\n"
                    "Industrial late-fusion deployment (continuous acoustic features + discrete grammar explanations), "
                    "C++ runtime optimization on ARM Cortex-A, and Modbus TCP / IEC 61850 streaming connectors for Schneider EcoStruxure Asset Advisor."
                )
                set_text(shape, next_text, font_size=10)

    # =========================================================================
    # SLIDE 5 (NEW): Scientific Rigor & Honest Engineering (Priority 3)
    # Cloned from Slide 3's two-card layout
    # =========================================================================
    print("Creating Slide 5: Scientific Rigor & Limitations...")
    s5_new = prs.slides.add_slide(prs.slide_layouts[0])
    for sp in list(s5_new.shapes):
        sp._element.getparent().remove(sp._element)
    for sp in s3.shapes:
        new_elem = copy.deepcopy(sp._element)
        s5_new.shapes._spTree.append(new_elem)

    for shape in s5_new.shapes:
        if shape.has_text_frame:
            t = shape.text_frame.text.strip()
            if "Innovation & Differentiation" in t:
                set_text(shape, "Scientific Rigor & Honest Engineering", font_size=20, bold=True, color=RGBColor(0, 160, 50))
            elif "What makes your approach genuinely new?" in t:
                set_text(shape, "What Works, What Does Not Yet Work Universally, and Architectural Responses", font_size=11)
            elif "INNOVATION & ORIGINALITY" in t:
                set_text(shape, "HONEST LIMITATIONS & DEFENSE", font_size=9, bold=True)
            elif "TODAY'S APPROACH" in t:
                set_text(shape, "WHAT WORKS (BENCHMARK VALIDATED)", font_size=12, bold=True, color=RGBColor(0, 140, 40))
            elif "YOUR APPROACH" in t:
                set_text(shape, "LIMITATIONS & ARCHITECTURAL RESPONSES", font_size=12, bold=True, color=RGBColor(180, 50, 0))
            elif shape.shape_id == 11:  # Left card body
                works_text = (
                    "✔ Explicit-Duration Semi-Markov Modeling:\n"
                    "  Discrete events + non-parametric dwell p_m(d) cut RUL error by 47–53% vs deep CNN/LSTM baselines (18.1 vs 34.0/38.9 RMSE on DS01).\n\n"
                    "✔ Stable Event Vocabulary Across Seeds:\n"
                    "  Vector Quantization (K=16) discovers repeatable operational tokens across 10 random seeds (0.94+ cosine stability, low perplexity variance).\n\n"
                    "✔ Rigorous Conformal Risk Bounds on DS01:\n"
                    "  Leave-One-Engine-Out (LOEO) calibration achieves 92% empirical coverage at 90% nominal level, eliminating window pseudoreplication.\n\n"
                    "✔ Intrinsic Explanation Faithfulness:\n"
                    "  Masking flagged event windows directly reduces anomaly score by 0.18–0.72 (6–27%), proving explanations reflect actual model drivers."
                )
                set_text(shape, works_text, font_size=10)
            elif shape.shape_id == 14:  # Right card body
                limits_text = (
                    "✖ Acoustic Discretization Penalty:\n"
                    "  Discrete VQ loses high-frequency info on MIMII fans (AUROC 0.535 vs 0.628 continuous).\n"
                    "  ↳ Response: Industrial Late-Fusion (continuous acoustic head + discrete grammar explanations).\n\n"
                    "✖ Weak Degradation Health Index:\n"
                    "  Health index Spearman correlation is weak (0.10–0.21) due to constant predictor artifact.\n"
                    "  ↳ Response: Use RUL conformal bounds for replacement; decouple health scoring.\n\n"
                    "✖ DS02 Small-Sample Conformal Limits:\n"
                    "  DS02 conformal coverage fluctuates (55–100%) due to small sample size (N=5 calibration engines).\n"
                    "  ↳ Response: Honestly flagged as unreliable for decision-making pending larger fleet N.\n\n"
                    "✖ Unproven Field Deployment:\n"
                    "  N-CMAPSS/MIMII are surrogate benchmarks; unvalidated on live Schneider switchgear.\n"
                    "  ↳ Response: Advisory decision support only (human-in-the-loop); pilot on Harmony iPC."
                )
                set_text(shape, limits_text, font_size=10)
            elif t == "5":
                set_text(shape, "5", font_size=10)

    # =========================================================================
    # SLIDE 6 (Old Slide 5): Meet the Team
    # =========================================================================
    s6_team = prs.slides[4]
    card_map = {
        10: ("Atharva Mendhulkar", 13, True),
        11: ("Computer Science & Engineering", 10, False),
        12: ("Lead AI/ML Architect & Backend — 1D Causal CNNs, Vector Quantizer (STE), Explicit-Duration HSMM Grammar, Automated Test Suite (260 tests), and Real-Time CLI Engine.", 9, False),
        16: ("Divyansh Seth", 13, True),
        17: ("Electrical & Electronics Engineering", 10, False),
        18: ("Power Systems & Asset Domain Lead — Switchgear Degradation Physics (Contact Resistance, Partial Discharge, Arc-Flash Risk), Sensor Telemetry, Conformal Safety Bounds & Business ROI Model.", 9, False),
        22: ("Academic Mentorship", 13, True),
        23: ("Dept. of Electrical & Computer Eng.", 10, False),
        24: ("Faculty Advisory — Guidance on stochastic processes, semi-Markov likelihood formulations, non-exponential wear distributions, and locked experimental peer review protocols.", 9, False),
        28: ("Schneider Electric Mentorship", 13, True),
        29: ("SE Electrifiers & Technical Mentors", 10, False),
        30: ("Industry Alignment — Mapped during Phase 3 Mentoring Cadence for EcoStruxure Asset Advisor integration, Masterpact MTZ telemetry contracts, and pilot testbench roadmaps.", 9, False),
    }
    for shape in s6_team.shapes:
        if shape.shape_id in card_map and shape.has_text_frame:
            val, sz, bld = card_map[shape.shape_id]
            set_text(shape, val, font_size=sz, bold=bld)
        elif shape.has_text_frame and shape.text_frame.text.strip().startswith("Faculty Mentor:"):
            set_text(shape, "Cross-Disciplinary Team: Computer Science + Electrical Engineering | Dedicated Faculty & Industry Mentor Cadence", font_size=11, bold=True)

    # =========================================================================
    # SLIDE 7 (Old Slide 6): Roadmap & Next Steps
    # =========================================================================
    s7_roadmap = prs.slides[5]
    for shape in s7_roadmap.shapes:
        if shape.has_text_frame:
            t = shape.text_frame.text.strip()
            if t.startswith("[ what you'll refine during the mentoring"):
                step1_text = (
                    "• Collaborate with assigned Schneider Electrifiers to validate switchgear sensor telemetry contracts.\n"
                    "• Calibrate domain adaptation on physical breaker thermal cycling and contact resistance datasets.\n"
                    "• Benchmark and optimize sub-15ms edge inference on ARM Cortex-A industrial hardware."
                )
                set_text(shape, step1_text, font_size=11)
            elif t.startswith("[ how you'd use the Schneider internship"):
                step2_text = (
                    "• Embed with Schneider R&D to port AEGIS onto Schneider Harmony industrial edge iPCs.\n"
                    "• Integrate event grammar telemetry streams into EcoStruxure Asset Advisor.\n"
                    "• Conduct live hardware-in-the-loop pilot testing on Masterpact MTZ circuit breaker test benches."
                )
                set_text(shape, step2_text, font_size=11)
            elif t.startswith("[ how this could scale across data centers"):
                step3_text = (
                    "• Scale across the global EcoStruxure fleet protecting millions of MV/LV electrical breakers and hyperscale data centers.\n"
                    "• Eliminate $1.08M downtime events per facility through conformal risk-bounded predictive replacement.\n"
                    "• Establish the industry benchmark for auditable, neuro-symbolic predictive intelligence."
                )
                set_text(shape, step3_text, font_size=11)

    # =========================================================================
    # SLIDE 8 (Old Slide 7): Thank You & Contact
    # =========================================================================
    s8_thanks = prs.slides[6]
    for shape in s8_thanks.shapes:
        if shape.has_text_frame:
            t = shape.text_frame.text.strip()
            if "[ Team Name ]" in t:
                set_text(shape, "AEGIS — Event Grammar Probabilistic Model (EGPM)  ·  Questions Welcome", font_size=18, bold=True, color=RGBColor(0, 160, 50))
            elif "[ Team leader contact email" in t:
                set_text(shape, "Atharva Mendhulkar (Atharvamendhulkar01@gmail.com) & Divyansh Seth\nProject Codebase & Verification: github.com/Atharva-Mendhulkar/aegis-1", font_size=13, bold=False)

    # =========================================================================
    # REORDER SLIDES SO SLIDE 5 IS INSERTED RIGHT AFTER SLIDE 4
    # Current order in prs: [S1, S2, S3, S4, S6_team, S7_road, S8_thanks, S5_new]
    # We want: [S1, S2, S3, S4, S5_new, S6_team, S7_road, S8_thanks]
    # =========================================================================
    sldIdLst = prs._element.xpath('//p:sldIdLst')[0]
    s5_elem = sldIdLst[-1]  # The newly appended slide 5
    sldIdLst.remove(s5_elem)
    sldIdLst.insert(4, s5_elem)  # Insert at index 4 (5th slide)

    # =========================================================================
    # SLIDE 9: MANDATORY Tagline Slide 13 from Kickoff PDF ("Advancing Energy Tech")
    # =========================================================================
    if IMG_SLIDE13.exists():
        print(f"Adding Mandatory Tagline Slide 9 from {IMG_SLIDE13.name}...")
        blank_layout = prs.slide_layouts[0]
        s9 = prs.slides.add_slide(blank_layout)
        for sp in list(s9.shapes):
            sp._element.getparent().remove(sp._element)
        s9.shapes.add_picture(str(IMG_SLIDE13), 0, 0, width=prs.slide_width, height=prs.slide_height)

    # =========================================================================
    # SLIDE 10: MANDATORY Closing CTA Slide 14 from Kickoff PDF ("se.com CTA")
    # =========================================================================
    if IMG_SLIDE14.exists():
        print(f"Adding Mandatory Closing CTA Slide 10 from {IMG_SLIDE14.name}...")
        blank_layout = prs.slide_layouts[0]
        s10 = prs.slides.add_slide(blank_layout)
        for sp in list(s10.shapes):
            sp._element.getparent().remove(sp._element)
        s10.shapes.add_picture(str(IMG_SLIDE14), 0, 0, width=prs.slide_width, height=prs.slide_height)

    prs.save(OUTPUT_PATH)
    print(f"\nSuccessfully generated complete 10-slide pitch deck: {OUTPUT_PATH.name} (Total Slides: {len(prs.slides)})")


if __name__ == "__main__":
    main()
